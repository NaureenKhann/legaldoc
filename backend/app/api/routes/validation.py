from typing import Any

import uuid
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse, Response
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import (
    CaseData as CaseDataModel,
    EvaluationReport as EvaluationReportModel,
    GenerationRun,
    ValidationRun as ValidationRunModel,
)
from app.db.session import get_db
from app.schemas.case import CaseData
from app.schemas.evaluation import EvaluationReportResponse
from app.schemas.generation import GenerationResult
from app.schemas.validation import ValidationRunResult
from app.services.reporting.reporter import EvaluationReporter
from app.services.scoring.engine import ScoringEngine
from app.services.validation.engine import ValidationEngine

router = APIRouter()
eval_router = APIRouter()


@router.post("/generation-runs/{run_id}/validate", response_model=ValidationRunResult)
async def validate_generation(
    run_id: str,
    db: AsyncSession = Depends(get_db),
) -> ValidationRunResult:
    run = await _get_run_or_404(run_id, db)

    if not run.output_document_path:
        raise HTTPException(status_code=422, detail="Generation not completed — no output document")

    # Load case data
    cd = await _get_case_data(str(run.matter_id), db)

    # Read generated document text (re-extract from DOCX for independent validation)
    doc_text = _read_docx_text(run.output_document_path)

    # Build minimal GenerationResult for validator
    meta = run.generation_metadata or {}
    gen_result = GenerationResult(
        matter_id=str(run.matter_id),
        run_id=run_id,
        file_path=run.output_document_path,
        generated_sections=meta.get("generated_sections", []),
        body_paragraph_count=meta.get("body_paragraph_count", 0),
        verification_range=meta.get("verification_range", ""),
        drafted_paragraphs=[],
        model_name=run.model_name or "",
        prompt_version=run.prompt_version or "v1.0",
    )

    # Run validation
    engine = ValidationEngine()
    val_result = engine.validate(cd, gen_result, doc_text)

    # Persist
    val_run = ValidationRunModel(
        id=uuid.uuid4(),
        generation_run_id=uuid.UUID(run_id),
        checks_json=[c.model_dump(mode="json") for c in val_result.checks],
        score_json=None,
        issue_count=val_result.failed,
    )
    db.add(val_run)
    await db.commit()
    await db.refresh(val_run)

    # Score and generate report
    scorer = ScoringEngine()
    score = scorer.score(val_result)

    reporter = EvaluationReporter()
    report = reporter.generate_report(
        matter_id=str(run.matter_id),
        generation_run_id=run_id,
        validation_run_id=str(val_run.id),
        validation_result=val_result,
        score=score,
    )

    # Save report
    rep_model = EvaluationReportModel(
        id=uuid.uuid4(),
        validation_run_id=val_run.id,
        report_markdown=report.report_markdown,
        report_json=report.model_dump(mode="json"),
    )
    db.add(rep_model)

    val_run.score_json = score.model_dump(mode="json")
    await db.commit()

    return val_result


@router.get("/generation-runs/{run_id}/validation", response_model=ValidationRunResult)
async def get_validation(
    run_id: str,
    db: AsyncSession = Depends(get_db),
) -> ValidationRunResult:
    result = await db.execute(
        select(ValidationRunModel).where(
            ValidationRunModel.generation_run_id == uuid.UUID(run_id)
        )
    )
    val_run = result.scalar_one_or_none()
    if not val_run:
        raise HTTPException(status_code=404, detail="Validation not yet run for this generation")

    from app.schemas.validation import ValidationCheck, CheckStatus, CheckSeverity, CheckCategory

    checks = [ValidationCheck.model_validate(c) for c in val_run.checks_json]
    passed = sum(1 for c in checks if c.status == CheckStatus.PASS)
    failed = sum(1 for c in checks if c.status == CheckStatus.FAIL)
    warnings = sum(1 for c in checks if c.status == CheckStatus.WARNING)

    return ValidationRunResult(
        generation_run_id=run_id,
        checks=checks,
        total_checks=len(checks),
        passed=passed,
        failed=failed,
        warnings=warnings,
        not_applicable=len(checks) - passed - failed - warnings,
        critical_failures=[
            c.check_id for c in checks
            if c.status == CheckStatus.FAIL and c.severity == CheckSeverity.ERROR
        ],
    )


@eval_router.get("/generation-runs/{run_id}/evaluation", response_model=EvaluationReportResponse)
async def get_evaluation(
    run_id: str,
    db: AsyncSession = Depends(get_db),
) -> EvaluationReportResponse:
    # Find validation run
    val_result = await db.execute(
        select(ValidationRunModel).where(
            ValidationRunModel.generation_run_id == uuid.UUID(run_id)
        )
    )
    val_run = val_result.scalar_one_or_none()
    if not val_run:
        raise HTTPException(status_code=404, detail="Run validation first")

    # Find evaluation report
    rep_result = await db.execute(
        select(EvaluationReportModel).where(
            EvaluationReportModel.validation_run_id == val_run.id
        )
    )
    rep = rep_result.scalar_one_or_none()
    if not rep:
        raise HTTPException(status_code=404, detail="Evaluation report not found")

    from app.schemas.evaluation import EvaluationReport
    report = EvaluationReport.model_validate(rep.report_json)

    return EvaluationReportResponse(
        matter_id=report.matter_id,
        generation_run_id=run_id,
        report=report,
        download_url=f"/api/v1/generation-runs/{run_id}/evaluation/download",
    )


@eval_router.get("/generation-runs/{run_id}/evaluation/download")
async def download_evaluation(
    run_id: str,
    format: str = "markdown",
    db: AsyncSession = Depends(get_db),
) -> Response:
    val_result = await db.execute(
        select(ValidationRunModel).where(
            ValidationRunModel.generation_run_id == uuid.UUID(run_id)
        )
    )
    val_run = val_result.scalar_one_or_none()
    if not val_run:
        raise HTTPException(status_code=404, detail="Validation run not found")

    rep_result = await db.execute(
        select(EvaluationReportModel).where(
            EvaluationReportModel.validation_run_id == val_run.id
        )
    )
    rep = rep_result.scalar_one_or_none()
    if not rep:
        raise HTTPException(status_code=404, detail="Report not found")

    if format == "json":
        import json
        return Response(
            content=json.dumps(rep.report_json, indent=2),
            media_type="application/json",
            headers={"Content-Disposition": f"attachment; filename=evaluation_{run_id}.json"},
        )

    if format == "docx":
        docx_bytes = _generate_evaluation_docx(rep.report_json)
        return Response(
            content=docx_bytes,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename=evaluation_{run_id}.docx"},
        )

    return Response(
        content=rep.report_markdown,
        media_type="text/markdown",
        headers={"Content-Disposition": f"attachment; filename=evaluation_{run_id}.md"},
    )


# ── Helpers ────────────────────────────────────────────────────────────────────

async def _get_run_or_404(run_id: str, db: AsyncSession) -> GenerationRun:
    result = await db.execute(
        select(GenerationRun).where(GenerationRun.id == uuid.UUID(run_id))
    )
    run = result.scalar_one_or_none()
    if not run:
        raise HTTPException(status_code=404, detail="Generation run not found")
    return run


async def _get_case_data(matter_id: str, db: AsyncSession) -> CaseData:
    result = await db.execute(
        select(CaseDataModel).where(CaseDataModel.matter_id == uuid.UUID(matter_id))
    )
    record = result.scalar_one_or_none()
    if not record:
        raise HTTPException(status_code=404, detail="Case data not found")
    return CaseData.model_validate(record.structured_json)


def _read_docx_text(file_path: str) -> str:
    """Re-extract text from generated DOCX for independent validation."""
    try:
        from docx import Document
        doc = Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception:
        return ""


def _generate_evaluation_docx(report_json: dict[str, Any]) -> bytes:
    """Generate a formal DOCX evaluation report."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import io

    doc = Document()

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("LegalDoc Intelligence Platform — Evaluation Report")
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(88, 28, 135)

    # Score Box
    score = report_json.get("score", {})
    p_score = doc.add_paragraph()
    p_score.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_score = p_score.add_run(f"\nOverall Score: {score.get('overall_score', 0):.1f} / 100  (Grade {score.get('grade', 'F')})\n")
    r_score.bold = True
    r_score.font.size = Pt(14)
    r_score.font.name = "Calibri"

    # Category Breakdown Table
    doc.add_heading("Category Breakdown", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Category"
    hdr_cells[1].text = "Score"
    hdr_cells[2].text = "Explanation"

    for cat in score.get("categories", []):
        row_cells = table.add_row().cells
        row_cells[0].text = str(cat.get("category", "")).replace("_", " ").title()
        row_cells[1].text = f"{cat.get('earned_points', 0):.1f} / {cat.get('max_points', 0)}"
        row_cells[2].text = str(cat.get("explanation", ""))

    # Issues Section
    issues = report_json.get("all_issues", [])
    if issues:
        doc.add_heading("Detected Audit Issues & Recommendations", level=2)
        for iss in issues:
            p_iss = doc.add_paragraph(style="List Bullet")
            r_id = p_iss.add_run(f"[{iss.get('status', '')}] {iss.get('check_id', '')}: ")
            r_id.bold = True
            p_iss.add_run(f"{iss.get('message', '')}\n")
            if iss.get("suggested_correction"):
                p_iss.add_run(f"Fix Suggestion: {iss.get('suggested_correction')}")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
