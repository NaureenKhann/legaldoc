"""
DOCX Generator — uses python-docx to produce a properly formatted Affidavit in Reply.
Python controls ALL structure. The LLM is NOT involved in document layout.
"""
from __future__ import annotations

import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import structlog
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement

from app.core.config import get_settings
from app.schemas.case import CaseData
from app.schemas.generation import DraftedParagraph, GenerationResult
from app.schemas.mapping import MappingPlan

logger = structlog.get_logger(__name__)
settings = get_settings()


@dataclass
class GenerationMetadata:
    file_path: str
    filename: str
    generated_sections: list[str] = field(default_factory=list)
    body_paragraph_count: int = 0
    verification_range: str = ""
    warnings: list[str] = field(default_factory=list)


class DocxGenerator:
    """
    Produces a legally formatted Affidavit in Reply DOCX.
    Python controls every structural element.
    """

    FONT_NAME = "Times New Roman"
    FONT_SIZE = Pt(12)
    HEADING_SIZE = Pt(12)
    MARGIN_INCHES = 1.0

    def __init__(self, storage_root: Path | None = None) -> None:
        self.storage_root = storage_root or settings.STORAGE_PATH
        self._doc: Document | None = None

    def generate(
        self,
        case_data: CaseData,
        mapping: MappingPlan,
        drafted_paragraphs: list[DraftedParagraph],
        matter_id: str,
        run_id: str,
        output_path: str | None = None,
    ) -> GenerationMetadata:
        log = logger.bind(matter_id=matter_id, run_id=run_id)
        log.info("docx_generation_start")

        self._doc = Document()
        self._set_document_margins()
        self._set_default_style()

        generated_sections: list[str] = []
        warnings: list[str] = []

        # ── 1. Forum Heading ─────────────────────────────────────────────────
        self.add_forum_heading(case_data.court.court_name)
        generated_sections.append("forum_heading")

        # ── 2. Jurisdiction ──────────────────────────────────────────────────
        self.add_jurisdiction_heading(case_data.court.jurisdiction)
        generated_sections.append("jurisdiction")

        # ── 3. Case Number ───────────────────────────────────────────────────
        self.add_case_number(
            case_data.court.proceeding_type,
            case_data.court.case_number,
            case_data.court.year,
        )
        generated_sections.append("case_number")

        # ── 4. Cause Title ───────────────────────────────────────────────────
        self.add_cause_title(case_data)
        generated_sections.append("cause_title")

        # ── 5. Affidavit Title ───────────────────────────────────────────────
        self.add_affidavit_title(case_data.answering_respondent_number)
        generated_sections.append("affidavit_title")

        # ── 6. Deponent Clause ───────────────────────────────────────────────
        self.add_deponent_clause(case_data)
        generated_sections.append("deponent_clause")

        # ── 7. Body Paragraphs ───────────────────────────────────────────────
        for dp in sorted(drafted_paragraphs, key=lambda p: p.paragraph_number):
            self.add_body_paragraph(dp.paragraph_number, dp.text)
            if dp.unsupported_claims:
                warnings.append(
                    f"Para {dp.paragraph_number}: possible unsupported claims: "
                    + "; ".join(dp.unsupported_claims)
                )
        generated_sections.append("body_paragraphs")

        body_count = len(drafted_paragraphs)

        # ── 8. Prayer ────────────────────────────────────────────────────────
        self.add_prayer(case_data.prayer.items)
        generated_sections.append("prayer")

        # ── 9. Jurat ─────────────────────────────────────────────────────────
        self.add_jurat(case_data.verification.place, case_data.verification.date)
        generated_sections.append("jurat")

        # ── 10. Verification ─────────────────────────────────────────────────
        self.add_verification(case_data, body_count)
        generated_sections.append("verification")

        # ── 11. Advocate Block ───────────────────────────────────────────────
        self.add_advocate_block(
            case_data.advocate_details.firm_name,
            case_data.advocate_details.acting_for,
        )
        generated_sections.append("advocate_block")

        # ── Save ─────────────────────────────────────────────────────────────
        if output_path:
            target_path = Path(output_path)
            target_path.parent.mkdir(parents=True, exist_ok=True)
            filename = target_path.name
        else:
            output_dir = self.storage_root / "generated" / matter_id
            output_dir.mkdir(parents=True, exist_ok=True)
            filename = f"affidavit_in_reply_{run_id}.docx"
            target_path = output_dir / filename

        self._doc.save(str(target_path))
        log.info("docx_saved", path=str(target_path), paragraphs=body_count)

        verification_range = (
            f"1 to {body_count}" if body_count > 1 else "1"
        )

        return GenerationMetadata(
            file_path=str(target_path),
            filename=filename,
            generated_sections=generated_sections,
            body_paragraph_count=body_count,
            verification_range=f"paragraphs {verification_range}",
            warnings=warnings,
        )

    # ── Builder Functions ─────────────────────────────────────────────────────

    def add_forum_heading(self, court_name: str) -> None:
        p = self._doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(court_name.upper())
        run.bold = True
        run.font.name = self.FONT_NAME
        run.font.size = self.HEADING_SIZE

    def add_jurisdiction_heading(self, jurisdiction: str) -> None:
        p = self._doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(jurisdiction.upper())
        run.bold = True
        run.font.name = self.FONT_NAME
        run.font.size = self.HEADING_SIZE

    def add_case_number(self, proceeding_type: str, case_number: str, year: int) -> None:
        self._doc.add_paragraph()  # spacing
        p = self._doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(
            f"{proceeding_type.upper()} NO. {case_number} OF {year}"
        )
        run.bold = True
        run.font.name = self.FONT_NAME
        run.font.size = self.FONT_SIZE

    def add_cause_title(self, case_data: CaseData) -> None:
        self._doc.add_paragraph()
        # Petitioner
        p = self._doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(case_data.petitioner.name.upper())
        r.font.name = self.FONT_NAME
        r.font.size = self.FONT_SIZE

        p2 = self._doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run("...PETITIONER")
        r2.font.name = self.FONT_NAME
        r2.font.size = self.FONT_SIZE

        # Versus
        pv = self._doc.add_paragraph()
        pv.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pv.add_run("VERSUS").font.bold = True

        # Respondents
        for r_party in sorted(case_data.respondents, key=lambda x: x.respondent_number or 99):
            p3 = self._doc.add_paragraph()
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run3 = p3.add_run(r_party.name.upper())
            run3.font.name = self.FONT_NAME
            run3.font.size = self.FONT_SIZE

            p4 = self._doc.add_paragraph()
            p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run4 = p4.add_run(f"...RESPONDENT NO. {r_party.respondent_number}")
            run4.font.name = self.FONT_NAME
            run4.font.size = self.FONT_SIZE

    def add_affidavit_title(self, respondent_number: int) -> None:
        self._doc.add_paragraph()
        p = self._doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(
            f"AFFIDAVIT IN REPLY ON BEHALF OF RESPONDENT NO. {respondent_number}"
        )
        run.bold = True
        run.font.name = self.FONT_NAME
        run.font.size = self.FONT_SIZE
        self._doc.add_paragraph()

    def add_deponent_clause(self, case_data: CaseData) -> None:
        d = case_data.deponent
        text = (
            f"I, {d.name}, {d.designation}, {d.organization}, "
            f"{d.address}, the deponent herein, acting for and on behalf of "
            f"{case_data.filed_on_behalf_of}, do hereby state as follows:"
        )
        p = self._doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = self.FONT_NAME
        run.font.size = self.FONT_SIZE
        self._doc.add_paragraph()

    def add_body_paragraph(self, number: int, text: str) -> None:
        p = self._doc.add_paragraph()
        run_num = p.add_run(f"{number}. ")
        run_num.bold = True
        run_num.font.name = self.FONT_NAME
        run_num.font.size = self.FONT_SIZE

        run_text = p.add_run(text)
        run_text.font.name = self.FONT_NAME
        run_text.font.size = self.FONT_SIZE

    def add_prayer(self, prayer_items: list[str]) -> None:
        self._doc.add_paragraph()
        p_label = self._doc.add_paragraph()
        r = p_label.add_run("PRAYER")
        r.bold = True
        r.font.name = self.FONT_NAME
        r.font.size = self.FONT_SIZE

        p_intro = self._doc.add_paragraph()
        p_intro.add_run(
            "It is, therefore, respectfully prayed that this Hon'ble Court may be "
            "pleased to:"
        ).font.name = self.FONT_NAME

        for i, item in enumerate(prayer_items):
            letter = chr(ord("a") + i)  # a, b, c, ... NOT 1, 2, 3
            p = self._doc.add_paragraph()
            run = p.add_run(f"({letter}) {item}")
            run.font.name = self.FONT_NAME
            run.font.size = self.FONT_SIZE

        self._doc.add_paragraph()

    def add_jurat(self, place: str, date: str) -> None:
        p = self._doc.add_paragraph()
        run = p.add_run(
            f"Solemnly affirmed at {place} on this {date}."
        )
        run.font.name = self.FONT_NAME
        run.font.size = self.FONT_SIZE

    def add_verification(self, case_data: CaseData, body_count: int) -> None:
        self._doc.add_paragraph()
        # Range is calculated from actual body_count — never copied from reference
        verification_range = f"1 to {body_count}" if body_count > 1 else "1"
        p = self._doc.add_paragraph()
        run = p.add_run(
            f"VERIFICATION: I, {case_data.deponent.name}, the deponent above-named, "
            f"do hereby verify that the contents of paragraphs {verification_range} of "
            f"this Affidavit in Reply are true and correct to the best of my knowledge, "
            f"information and belief. No part of it is false and nothing material has been "
            f"concealed therefrom.\n\n"
            f"Verified at {case_data.verification.place} on "
            f"{case_data.verification.date}."
        )
        run.font.name = self.FONT_NAME
        run.font.size = self.FONT_SIZE

    def add_advocate_block(self, firm_name: str, acting_for: str) -> None:
        self._doc.add_paragraph()
        p = self._doc.add_paragraph()
        run = p.add_run(
            f"{firm_name}\nAdvocates for {acting_for}"
        )
        run.font.name = self.FONT_NAME
        run.font.size = self.FONT_SIZE

    # ── Document Setup ────────────────────────────────────────────────────────

    def _set_document_margins(self) -> None:
        from docx.shared import Inches
        for section in self._doc.sections:
            section.top_margin = Inches(self.MARGIN_INCHES)
            section.bottom_margin = Inches(self.MARGIN_INCHES)
            section.left_margin = Inches(1.25)
            section.right_margin = Inches(1.0)

    def _set_default_style(self) -> None:
        style = self._doc.styles["Normal"]
        font = style.font
        font.name = self.FONT_NAME
        font.size = self.FONT_SIZE
